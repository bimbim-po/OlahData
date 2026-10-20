import os
import sys
import json
import io
import pytest
import docx
import pypdf

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from exporter.docx_export import generate_docx_report
from exporter.pdf_export import generate_pdf_report
from guard import NumberGuard, NumberGuardViolationError

@pytest.fixture
def golden_results():
    fixture_path = os.path.join(os.path.dirname(__file__), "..", "..", "..", "fixtures", "hasil_acuan_bersih.json")
    with open(fixture_path, "r", encoding="utf-8") as f:
        return json.load(f)

def test_chapter4_docx_export_structure_and_margins(golden_results):
    """
    Verifies that the generated Word (.docx) document strictly adheres to:
    - 4-3-3-3 cm margins (Left 4 cm, Top 3 cm, Right 3 cm, Bottom 3 cm)
    - Full Table Matrix 4.1 to 4.15
    - Editable native tables with APA 7th style
    - List of Tables (Daftar Tabel)
    - Non-hallucinatory discussion scaffolds
    - Verifiable audit trail and Signed Academic Declaration
    """
    docx_bytes = generate_docx_report(
        golden_results,
        project_title="Pengaruh Kualitas Layanan dan Kepuasan terhadap Loyalitas Pelanggan",
        include_declaration=True,
        watermark=False
    )
    assert len(docx_bytes) > 20000, "DOCX file size should be substantial"

    doc = docx.Document(io.BytesIO(docx_bytes))
    sec = doc.sections[0]

    # Verify Margins
    assert round(sec.left_margin.cm, 1) == 4.0, f"Expected left margin 4.0 cm, got {sec.left_margin.cm}"
    assert round(sec.top_margin.cm, 1) == 3.0, f"Expected top margin 3.0 cm, got {sec.top_margin.cm}"
    assert round(sec.right_margin.cm, 1) == 3.0, f"Expected right margin 3.0 cm, got {sec.right_margin.cm}"
    assert round(sec.bottom_margin.cm, 1) == 3.0, f"Expected bottom margin 3.0 cm, got {sec.bottom_margin.cm}"

    # Verify Table Count (all tables present)
    assert len(doc.tables) >= 15, f"Expected at least 15 tables, found {len(doc.tables)}"

    # Collect all text
    full_text = []
    for p in doc.paragraphs:
        full_text.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                full_text.append(cell.text)
    doc_str = " ".join(full_text)

    # 1. Exact Numerical Match with golden fixture
    assert "0,4710" in doc_str, "KL1 rHitung 0,4710 must match golden reference"
    assert "-0,0471" in doc_str, "KL5 rHitung -0,0471 must match golden reference"
    assert "Gugur (Tidak Valid)" in doc_str, "KL5 must be marked as Gugur (Tidak Valid)"
    assert "0,3863" in doc_str, "R-Square 0,3863 must match golden reference"
    assert "30,5246" in doc_str, "F-Hitung 30,5246 must match golden reference"
    assert "0,5546" in doc_str, "Lilliefors Sig 0,5546 must match golden reference"
    assert "0,7331" in doc_str, "Shapiro-Wilk Sig 0,7331 must match golden reference"
    assert "1,0514" in doc_str, "VIF 1,0514 must match golden reference"
    assert "0,9511" in doc_str, "Tolerance 0,9511 must match golden reference"

    # 2. Academic Discussion Scaffold
    assert "[Tambahkan teori dan penelitian terdahulu Anda di sini]" in doc_str, "Discussion scaffold placeholder must be present"
    assert "Apakah temuan ini sejalan atau bertentangan dengan penelitian terdahulu?" in doc_str, "Guiding questions must be present"

    # 3. Appendices
    assert "LAMPIRAN 1: JEJAK HITUNG & REPRODUKSIBILITAS DETERMINISTIK" in doc_str, "Audit trail appendix must be present"
    assert "LAMPIRAN 2: PERNYATAAN PENGGUNAAN ALAT BANTU AKADEMIK" in doc_str, "Academic declaration appendix must be present"
    assert "DAFTAR TABEL BAB IV" in doc_str, "List of tables must be present"

def test_chapter4_pdf_export_structure(golden_results):
    """
    Verifies that the generated PDF report is valid, contains multiple pages,
    and includes essential headers, tables, and declaration.
    """
    pdf_bytes = generate_pdf_report(
        golden_results,
        project_title="Pengaruh Kualitas Layanan dan Kepuasan terhadap Loyalitas Pelanggan",
        include_declaration=True,
        watermark=True
    )
    assert len(pdf_bytes) > 5000, "PDF file size should be substantial"

    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 4, f"Expected at least 4 pages in Chapter 4 PDF, got {len(reader.pages)}"

    pdf_text = ""
    for page in reader.pages:
        pdf_text += page.extract_text() or ""

    assert "BAB IV" in pdf_text, "Title BAB IV must be in PDF"
    assert "Tabel 4." in pdf_text, "Tables 4.x must be present in PDF"
    assert "0,4710" in pdf_text or "0.4710" in pdf_text, "KL1 rHitung must be present in PDF"
    assert "SURAT PERNYATAAN PENGGUNAAN ALAT BANTU" in pdf_text or "PERNYATAAN PENGGUNAAN ALAT" in pdf_text, "Declaration must be in PDF"

def test_number_guard_rejects_hallucinations(golden_results):
    """
    Verifies that Number Guard strictly catches and rejects any narrative text
    containing fabricated statistical values or wrong categories.
    """
    guard = NumberGuard(tolerance=0.01)

    # Valid narrative based on golden results (R-Square = 0.3863, F = 30.52)
    valid_narrative = (
        "Berdasarkan hasil analisis, nilai R-Square sebesar 0,386 dan F-Hitung sebesar 30,52. "
        "Seluruh butir konstruk KL terbukti andal dengan kategori Tinggi."
    )
    res_valid = guard.verify_narrative(valid_narrative, golden_results, strict=False)
    assert res_valid["valid"] is True, f"Valid narrative failed: {res_valid}"

    # Invalid narrative with fabricated R-Square (claims 0.850 instead of 0.386)
    invalid_r2_narrative = "Hasil pengujian menunjukkan nilai R-Square sebesar 0,850 yang sangat tinggi."
    with pytest.raises(NumberGuardViolationError) as exc_info:
        guard.verify_narrative(invalid_r2_narrative, golden_results, strict=True)
    assert "R-Square" in str(exc_info.value), "Should reject fabricated R-Square"

    # Invalid narrative with fabricated F-Hitung
    invalid_f_narrative = "Uji simultan menghasilkan F-Hitung sebesar 99,50 dengan signifikansi tinggi."
    with pytest.raises(NumberGuardViolationError) as exc_info:
        guard.verify_narrative(invalid_f_narrative, golden_results, strict=True)
    assert "F-Hitung" in str(exc_info.value), "Should reject fabricated F-Hitung"

    # Invalid narrative with wrong category (claims Rendah when true is Tinggi)
    invalid_kat_narrative = "Variabel KL berada pada kategori Rendah menurut penilaian responden."
    with pytest.raises(NumberGuardViolationError) as exc_info:
        guard.verify_narrative(invalid_kat_narrative, golden_results, strict=True)
    assert "Kategori" in str(exc_info.value), "Should reject incorrect category"
